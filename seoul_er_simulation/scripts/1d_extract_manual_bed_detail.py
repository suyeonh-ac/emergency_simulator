"""
1d단계: 사용자가 응급똑똑 앱에서 직접 확인해 정리한 병상/장비 데이터를 구조화

입력: data/manual_bed_detail.xlsx (사용자 제공 — 서울 법정 응급의료기관 51곳
      〈권역응급의료센터 7 + 지역응급의료센터 24 + 지역응급의료기관 20〉의
      응급실/중환자실/감염병상/응급전용병상/외상전용병상/입원실/기타병상/보유장비
      "전체(분모, 최대 수용치)" 실측)
출력: data/bed_detail_raw.json  (hpid로 매핑, 원본 컬럼 구조 그대로 보관)

주의: 이 파일은 74곳 중 51곳만 다룬다 — 응급실운영신고기관 23곳은 응급똑똑 앱에
세부 병상 정보가 없는 경우가 많아(소규모 병원) 대상에서 제외됨. 이 23곳은
scripts/1c_apply_egen_data.py가 만든 hperyn 기반 값을 그대로 쓴다
(scripts/1e_apply_manual_bed_detail.py에서 처리).
"""
import json
import os
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(HERE, '..', 'data')

SRC_XLSX = os.path.join(DATA_DIR, 'manual_bed_detail.xlsx')
HOSPITALS_JSON = os.path.join(DATA_DIR, 'hospitals.json')
OUT_JSON = os.path.join(DATA_DIR, 'bed_detail_raw.json')

# 병원명 표기 차이(법인명 접두어 등) 보정용 — norm()으로 양쪽을 정규화해 비교
STRIP_PREFIXES = [
    '학교법인', '재단법인', '의료법인', '사회복지법인', '의료재단', '복지재단',
    '고려중앙학원', '가톨릭학원', '한국보훈복지의료공단',
]


def norm(s):
    s = s or ''
    for pfx in STRIP_PREFIXES:
        s = s.replace(pfx, '')
    return s.replace(' ', '')


def main():
    hospitals = json.load(open(HOSPITALS_JSON, encoding='utf-8'))
    name_to_hpid = {norm(h['name']): h['hpid'] for h in hospitals}

    wb = load_workbook(SRC_XLSX, data_only=True)
    ws = wb['응급실_병상데이터']

    group_row = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_row = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
    groups = []
    last = None
    for g in group_row:
        if g is not None:
            last = g
        groups.append(last)

    result = {}
    unmatched = []
    for r in range(3, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if vals[1] is None:
            continue
        name = vals[1]
        hpid = name_to_hpid.get(norm(name))
        if hpid is None:
            unmatched.append(name)
            continue

        entry = {'name': name, 'grade_in_file': vals[2], 'address_in_file': vals[3], 'groups': {}}
        for i in range(4, len(vals)):
            group = groups[i]
            col = col_row[i]
            v = vals[i]
            entry['groups'].setdefault(group, {})[col] = v
        result[hpid] = entry

    print(f'{len(result)}개 병원 매칭, {len(unmatched)}개 미매칭: {unmatched}')

    with open(OUT_JSON, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print(f'{OUT_JSON} 작성 완료')


if __name__ == '__main__':
    main()
